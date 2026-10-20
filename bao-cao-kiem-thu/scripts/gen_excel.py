# -*- coding: utf-8 -*-
"""Sinh file Excel 125 test case tu tc_data.py"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tc_data import NHOM, AUTH, PROD, ORD, VOUC, REV, CHAT, AI, HR, ADM

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

GROUPS = [("AUTH",AUTH),("PROD",PROD),("ORD",ORD),("VOUC",VOUC),("REV",REV),
          ("CHAT",CHAT),("AI",AI),("HR",HR),("ADM",ADM)]

LOAI = {"P":"Positive","N":"Negative","B":"Boundary","S":"Security"}

HEADERS = ["Mã TC","Nhóm chức năng","UC liên quan","Tiêu đề kiểm thử",
           "Điều kiện tiên quyết","Dữ liệu đầu vào","Các bước thực hiện",
           "Kết quả mong đợi","Loại KT","Mức ưu tiên","Phương pháp",
           "Kết quả thực tế","Trạng thái","Người thực hiện","Ngày thực hiện","Ghi chú"]
WIDTHS  = [11, 24, 12, 34, 28, 30, 38, 46, 11, 12, 12, 26, 14, 14, 13, 20]

HDR_FILL  = PatternFill("solid", fgColor="1F3864")
HDR_FONT  = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
CELL_FONT = Font(name="Times New Roman", size=11)
TITLE_FONT= Font(name="Times New Roman", size=14, bold=True, color="1F3864")
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)

FILL_TYPE = {"P": "E2EFDA", "N": "FCE4D6", "B": "FFF2CC", "S": "FADBD8"}

DV_STATUS = '"Pass,Fail,Blocked,Chưa thực hiện"'


def style_header(ws):
    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c-1]
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "E2"


def write_rows(ws, rows, start=2):
    r = start
    for (ma, uc, tieu_de, tien_de, du_lieu, buoc, ket_qua, loai, uu_tien, pp) in rows:
        prefix = ma.split("-")[0]
        vals = [ma, NHOM[prefix], uc, tieu_de, tien_de, du_lieu, buoc, ket_qua,
                LOAI[loai], uu_tien, pp, "", "Chưa thực hiện", "", "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = CENTER if c in (1, 3, 9, 10, 11, 13, 15) else TOP_WRAP
        ws.cell(row=r, column=9).fill = PatternFill("solid", fgColor=FILL_TYPE[loai])
        r += 1
    # dropdown cho cot Trang thai
    dv = DataValidation(type="list", formula1=DV_STATUS, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{r-1}")
    ws.auto_filter.ref = f"A1:P{r-1}"
    return r - 1


def build():
    wb = Workbook()

    # ---------- Sheet tong hop ----------
    s = wb.active
    s.title = "Tổng hợp"
    s["A1"] = "TỔNG HỢP KẾT QUẢ KIỂM THỬ - HỆ THỐNG TIGER SHOP"
    s["A1"].font = TITLE_FONT
    s.merge_cells("A1:G1")
    s["A2"] = "Học phần: Đánh giá và Kiểm định chất lượng phần mềm"
    s["A2"].font = Font(name="Times New Roman", size=11, italic=True)
    s.merge_cells("A2:G2")

    hdr = ["STT","Mã nhóm","Nhóm chức năng","Tổng số TC","Pass","Fail","Not Run","Tỉ lệ Pass"]
    for c, h in enumerate(hdr, 1):
        cell = s.cell(row=4, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    for c, w in enumerate([6, 11, 30, 12, 9, 9, 11, 12], 1):
        s.column_dimensions[get_column_letter(c)].width = w

    r = 5
    for i, (key, rows) in enumerate(GROUPS, 1):
        n = len(rows)
        sheet = f"'{key}'"
        vals = [i, key, NHOM[key], n,
                f'=COUNTIF({sheet}!M:M,"Pass")',
                f'=COUNTIF({sheet}!M:M,"Fail")',
                f'=COUNTIF({sheet}!M:M,"Chưa thực hiện")+COUNTIF({sheet}!M:M,"Blocked")',
                f'=IF(D{r}=0,0,E{r}/D{r})']
        for c, v in enumerate(vals, 1):
            cell = s.cell(row=r, column=c, value=v)
            cell.font, cell.border = CELL_FONT, BORDER
            cell.alignment = TOP_WRAP if c == 3 else CENTER
        s.cell(row=r, column=8).number_format = "0.0%"
        r += 1

    tot = r
    s.cell(row=tot, column=1, value="")
    s.cell(row=tot, column=3, value="TỔNG CỘNG").font = Font(name="Times New Roman", size=11, bold=True)
    for c, col in zip(range(4, 8), "DEFG"):
        s.cell(row=tot, column=c, value=f"=SUM({col}5:{col}{tot-1})")
    s.cell(row=tot, column=8, value=f"=IF(D{tot}=0,0,E{tot}/D{tot})").number_format = "0.0%"
    for c in range(1, 9):
        cell = s.cell(row=tot, column=c)
        cell.font = Font(name="Times New Roman", size=11, bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E2F3")
        cell.border, cell.alignment = BORDER, CENTER

    # Thong ke theo loai kiem thu
    r2 = tot + 3
    s.cell(row=r2, column=1, value="PHÂN BỐ THEO LOẠI KIỂM THỬ").font = TITLE_FONT
    s.merge_cells(start_row=r2, start_column=1, end_row=r2, end_column=4)
    for c, h in enumerate(["Loại kiểm thử", "Số lượng", "Tỉ lệ", ""], 1):
        if h:
            cell = s.cell(row=r2+1, column=c, value=h)
            cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    all_rows = [row for _, rows in GROUPS for row in rows]
    total = len(all_rows)
    rr = r2 + 2
    for code, name in LOAI.items():
        n = sum(1 for x in all_rows if x[7] == code)
        for c, v in enumerate([name, n, n/total], 1):
            cell = s.cell(row=rr, column=c, value=v)
            cell.font, cell.border, cell.alignment = CELL_FONT, BORDER, CENTER
        s.cell(row=rr, column=3).number_format = "0.0%"
        s.cell(row=rr, column=1).fill = PatternFill("solid", fgColor=FILL_TYPE[code])
        rr += 1

    # Thong ke theo phuong phap
    r3 = rr + 2
    s.cell(row=r3, column=1, value="PHÂN BỐ THEO PHƯƠNG PHÁP THỰC HIỆN").font = TITLE_FONT
    s.merge_cells(start_row=r3, start_column=1, end_row=r3, end_column=4)
    for c, h in enumerate(["Phương pháp", "Số lượng", "Tỉ lệ"], 1):
        cell = s.cell(row=r3+1, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    rr = r3 + 2
    for pp in ["Tự động", "Thủ công"]:
        n = sum(1 for x in all_rows if x[9] == pp)
        for c, v in enumerate([pp, n, n/total], 1):
            cell = s.cell(row=rr, column=c, value=v)
            cell.font, cell.border, cell.alignment = CELL_FONT, BORDER, CENTER
        s.cell(row=rr, column=3).number_format = "0.0%"
        rr += 1

    # ---------- Sheet toan bo ----------
    ws_all = wb.create_sheet("Toàn bộ Test Case")
    style_header(ws_all)
    write_rows(ws_all, all_rows)

    # ---------- Sheet theo nhom ----------
    for key, rows in GROUPS:
        ws = wb.create_sheet(key)
        style_header(ws)
        write_rows(ws, rows)

    # ---------- Sheet truy vet ----------
    tv = wb.create_sheet("Truy vết UC - TC")
    tv["A1"] = "MA TRẬN TRUY VẾT GIỮA USE CASE VÀ TEST CASE"
    tv["A1"].font = TITLE_FONT
    tv.merge_cells("A1:D1")
    for c, h in enumerate(["Mã UC", "Số test case", "Danh sách mã test case", "Ghi chú"], 1):
        cell = tv.cell(row=3, column=c, value=h)
        cell.fill, cell.font, cell.alignment, cell.border = HDR_FILL, HDR_FONT, CENTER, BORDER
    for c, w in enumerate([10, 14, 76, 26], 1):
        tv.column_dimensions[get_column_letter(c)].width = w

    mapping = {}
    for row in all_rows:
        mapping.setdefault(row[1], []).append(row[0])
    r = 4
    for i in range(1, 26):
        uc = f"UC{i:02d}"
        tcs = mapping.get(uc, [])
        note = "" if tcs else "Chưa có test case trực tiếp, được bao phủ gián tiếp"
        for c, v in enumerate([uc, len(tcs), ", ".join(tcs), note], 1):
            cell = tv.cell(row=r, column=c, value=v)
            cell.font, cell.border = CELL_FONT, BORDER
            cell.alignment = CENTER if c in (1, 2) else TOP_WRAP
        if not tcs:
            tv.cell(row=r, column=1).fill = PatternFill("solid", fgColor="FCE4D6")
        r += 1

    out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "Tiger_Shop_TestCases.xlsx")
    wb.save(out)
    print("Đã tạo:", out)
    print("Tổng số test case:", total)
    chua = [f"UC{i:02d}" for i in range(1, 26) if f"UC{i:02d}" not in mapping]
    print("Use Case chưa có test case trực tiếp:", chua if chua else "không có")


if __name__ == "__main__":
    build()
