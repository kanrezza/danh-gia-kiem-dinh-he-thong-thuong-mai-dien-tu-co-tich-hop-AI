# -*- coding: utf-8 -*-
"""
Doi chieu ket qua Jest voi trang thai test case trong Tiger_Shop_TestCases.xlsx.

Tu ngay 28/09/2026, trang thai Pass, Fail cua tung test case duoc danh gia theo dac ta
o Chuong I va bon trang thai cua Bang 2.25, khong suy ra tu viec kich ban Jest chay xanh
hay do. Nhieu kich ban duoc viet de khang dinh hanh vi hien tai cua he thong (vi du
ADM-005b khang dinh mat khau bi dat ve hang so), nen kich ban chay xanh khong co nghia
la test case dat. Vi vay script nay KHONG ghi cot Trang thai va cot Ket qua thuc te.
Script chi lam ba viec:
  1. Doi chieu tung test case trong Excel voi cac kich ban Jest mang ma cua no
     (ke ca kich ban hau to b, c nhu AUTH-019b) va in danh sach can xu ly;
  2. Ghi bao cao doi chieu ra tep doi_chieu_jest.txt canh tep Excel;
  3. Khi chay voi --ghi-ngay: ghi cot Ngay thuc hien bang ngay chay Jest (startTime
     trong jest-results.json, gio Viet Nam) cho cac test case da danh gia, thuc hien
     bang Jest va co moi kich ban lien quan deu thanh cong. Khong o nao khac bi sua.

Chay:
  cd backend && npx jest --runInBand --json --outputFile=../bao-cao-kiem-thu/jest-results.json
  cd bao-cao-kiem-thu/scripts && python3 import_results.py              # chi doi chieu
  cd bao-cao-kiem-thu/scripts && python3 import_results.py --ghi-ngay   # doi chieu va ghi ngay

Ban cu (ghi de trang thai theo mau xanh do cua Jest) duoc giu o import_results_ban_cu.py.
"""
import argparse
import json
import os
import re
from datetime import datetime, timezone, timedelta

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ICT = timezone(timedelta(hours=7))

SHEET_CHINH = "Toàn bộ Test Case"
SHEET_NHOM = ["AUTH", "PROD", "ORD", "VOUC", "REV", "CHAT", "AI", "HR", "ADM"]
NGUOI_JEST = "Jest + Supertest"
DA_DANH_GIA = ("Pass", "Fail", "Blocked")
CHUA_THUC_HIEN = "Chưa thực hiện"

TC_RE = re.compile(r"\b(AUTH|PROD|ORD|VOUC|REV|CHAT|AI|HR|ADM)-(\d{3})([a-z]?)\b")

# Test case co bang chung tu kich ban khong mang ma test case (Bang 3.6)
BANG_CHUNG_KHAC = {
    "ORD-022": re.compile(r"STAT: nhan vien cap nhat trang thai don"),
}


def doc_jest(path):
    data = json.load(open(path, encoding="utf-8"))
    bat_dau = datetime.fromtimestamp(data["startTime"] / 1000, ICT)
    kich_ban = {}
    dem = {"passed": 0, "failed": 0, "khac": 0, "co_ma": 0}
    for suite in data["testResults"]:
        for t in suite.get("assertionResults", []):
            st, ten = t["status"], t["fullName"]
            dem[st if st in ("passed", "failed") else "khac"] += 1
            ma = {f"{p}-{n}" for p, n, _ in TC_RE.findall(ten)}
            if ma:
                dem["co_ma"] += 1
            for code, mau in BANG_CHUNG_KHAC.items():
                if mau.search(ten):
                    ma.add(code)
            for code in ma:
                kich_ban.setdefault(code, []).append((ten, st))
    return bat_dau, kich_ban, dem


def cot(ws):
    return {c.value: c.column for c in ws[1] if c.value}


def doc_trang_thai(ws):
    c = cot(ws)
    kq = {}
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, c["Mã TC"]).value
        if code:
            kq[code] = {
                "trang_thai": ws.cell(r, c["Trạng thái"]).value,
                "nguoi": ws.cell(r, c["Người thực hiện"]).value,
                "ngay": ws.cell(r, c["Ngày thực hiện"]).value,
            }
    return kq


def phan_loai(tc, kich_ban):
    nhom = {"danh_gia_lai": [], "kich_ban_moi": [], "thieu_kich_ban": [], "fail_con_nguyen": [],
            "khop": []}
    for code, x in tc.items():
        kb = kich_ban.get(code, [])
        loi = [ten for ten, st in kb if st != "passed"]
        if loi:
            nhom["danh_gia_lai"].append((code, x, "kịch bản không thành công: " + "; ".join(loi)))
        elif x["trang_thai"] == CHUA_THUC_HIEN and kb:
            nhom["kich_ban_moi"].append((code, x, "đã có kịch bản: " + "; ".join(t for t, _ in kb)))
        elif x["trang_thai"] in DA_DANH_GIA and x["nguoi"] == NGUOI_JEST and not kb:
            nhom["thieu_kich_ban"].append((code, x, "Excel ghi thực hiện bằng Jest nhưng lượt chạy không có kịch bản mang mã này"))
        elif x["trang_thai"] == "Fail" and kb:
            nhom["fail_con_nguyen"].append((code, x, "kịch bản vẫn khẳng định hành vi hiện tại, khiếm khuyết chưa được sửa"))
        else:
            nhom["khop"].append((code, x, ""))
    return nhom


def main():
    ap = argparse.ArgumentParser(description="Đối chiếu kết quả Jest với Tiger_Shop_TestCases.xlsx")
    ap.add_argument("ket_qua", nargs="?", default=os.path.join(BASE, "jest-results.json"))
    ap.add_argument("--xlsx", default=os.path.join(BASE, "Tiger_Shop_TestCases.xlsx"))
    ap.add_argument("--bao-cao", default=None, help="tệp báo cáo, mặc định doi_chieu_jest.txt cạnh tệp Excel")
    ap.add_argument("--ghi-ngay", action="store_true",
                    help="ghi cột Ngày thực hiện bằng ngày chạy Jest cho các test case không có gì thay đổi")
    a = ap.parse_args()
    bao_cao = a.bao_cao or os.path.join(os.path.dirname(os.path.abspath(a.xlsx)), "doi_chieu_jest.txt")

    if not os.path.exists(a.ket_qua):
        raise SystemExit(f"Không tìm thấy {a.ket_qua}. Hãy chạy Jest với --json trước.")

    bat_dau, kich_ban, dem = doc_jest(a.ket_qua)
    wb = openpyxl.load_workbook(a.xlsx)
    tc = doc_trang_thai(wb[SHEET_CHINH])
    nhom = phan_loai(tc, kich_ban)
    ngay = bat_dau.strftime("%d/%m/%Y")

    # Kiem tra sheet nhom co cung trang thai voi sheet chinh
    lech_sheet = []
    for name in SHEET_NHOM:
        if name in wb.sheetnames:
            for code, x in doc_trang_thai(wb[name]).items():
                if code not in tc or x["trang_thai"] != tc[code]["trang_thai"]:
                    lech_sheet.append(f"{name}: {code}")
    ma_la = sorted(set(kich_ban) - set(tc))

    dem_tt = {}
    for x in tc.values():
        dem_tt[x["trang_thai"]] = dem_tt.get(x["trang_thai"], 0) + 1

    dong = []
    dong.append(f"ĐỐI CHIẾU KẾT QUẢ JEST VỚI {os.path.basename(a.xlsx)}")
    dong.append(f"Lượt chạy Jest: {bat_dau.strftime('%d/%m/%Y %H:%M')} (giờ Việt Nam)")
    tong_kb = dem["passed"] + dem["failed"] + dem["khac"]
    dong.append(f"Số kịch bản: {tong_kb}; thành công {dem['passed']}; không thành công {dem['failed']}; khác {dem['khac']}")
    dong.append(f"Kịch bản mang mã test case (kể cả hậu tố b, c): {dem['co_ma']}")
    dong.append("Trạng thái trong Excel (giữ nguyên, không bị script sửa): "
                + "; ".join(f"{k} {v}" for k, v in sorted(dem_tt.items(), key=lambda kv: str(kv[0]))))
    dong.append("")

    muc = [
        ("danh_gia_lai", "CẦN ĐÁNH GIÁ LẠI THEO ĐẶC TẢ (hành vi hệ thống đã thay đổi)"),
        ("kich_ban_moi", "CÓ KỊCH BẢN MỚI, CẦN ĐÁNH GIÁ RỒI GHI TRẠNG THÁI THỦ CÔNG"),
        ("thieu_kich_ban", "THIẾU KỊCH BẢN TRONG LƯỢT CHẠY"),
        ("fail_con_nguyen", "FAIL THEO ĐẶC TẢ, KHIẾM KHUYẾT CHƯA ĐƯỢC SỬA (không cần xử lý)"),
    ]
    can_xu_ly = 0
    for key, tieu_de in muc:
        ds = nhom[key]
        dong.append(f"{tieu_de}: {len(ds)}")
        for code, x, ly_do in ds:
            dong.append(f"  {code} [{x['trang_thai']}] {ly_do}")
        if key in ("danh_gia_lai", "kich_ban_moi", "thieu_kich_ban"):
            can_xu_ly += len(ds)
        dong.append("")
    if lech_sheet:
        dong.append(f"SHEET NHÓM LỆCH TRẠNG THÁI VỚI SHEET {SHEET_CHINH}: " + ", ".join(lech_sheet))
        can_xu_ly += len(lech_sheet)
    if ma_la:
        dong.append("MÃ TRONG KỊCH BẢN KHÔNG CÓ TRONG EXCEL: " + ", ".join(ma_la))
        can_xu_ly += len(ma_la)

    # Ghi ngay cho cac test case khong co gi thay doi
    so_o = 0
    if a.ghi_ngay:
        du_dieu_kien = {code for code, x, _ in nhom["khop"] + nhom["fail_con_nguyen"]
                        if x["trang_thai"] in DA_DANH_GIA and x["nguoi"] == NGUOI_JEST and kich_ban.get(code)}
        for name in [SHEET_CHINH] + SHEET_NHOM:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            c = cot(ws)
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, c["Mã TC"]).value in du_dieu_kien:
                    o = ws.cell(r, c["Ngày thực hiện"])
                    if o.value != ngay:
                        o.value = ngay
                        so_o += 1
        if so_o:
            wb.calculation.fullCalcOnLoad = True
            wb.save(a.xlsx)
        dong.append(f"Đã ghi ngày {ngay} vào {so_o} ô Ngày thực hiện "
                    f"({len(du_dieu_kien)} test case). Cột Trạng thái và Kết quả thực tế không bị sửa.")
    else:
        dong.append("Chế độ chỉ đối chiếu: tệp Excel không bị sửa. Thêm --ghi-ngay để ghi ngày chạy Jest.")

    dong.append("KẾT LUẬN: " + ("không có test case nào cần xử lý." if can_xu_ly == 0
                                 else f"có {can_xu_ly} mục cần xử lý, xem các phần ở trên."))

    noi_dung = "\n".join(dong) + "\n"
    print(noi_dung, end="")
    with open(bao_cao, "w", encoding="utf-8") as f:
        f.write(noi_dung)
    print(f"Đã ghi báo cáo: {bao_cao}")


if __name__ == "__main__":
    main()
